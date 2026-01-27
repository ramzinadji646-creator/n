import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

class ReportService:
    @staticmethod
    def generate_recipes_excel(recipes):
        wb = Workbook()
        ws = wb.active
        ws.title = "Recipes"

        headers = ["ID", "Name", "Labor Hours", "Packaging Cost", "Selling Price", "Ingredients Count"]
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for r in recipes:
            ws.append([r.id, r.name, r.labor_hours, r.packaging_cost, r.selling_price, len(r.ingredients)])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_weekly_pdf(stats):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph("Weekly Bakery Report", styles['Title']))
        elements.append(Paragraph(f"Revenue: {stats['revenue']} DA", styles['Normal']))
        elements.append(Paragraph(f"Profit: {stats['profit']} DA", styles['Normal']))

        # Add table
        data = [["Product", "Quantity", "Revenue"]]
        for p in stats['products']:
            data.append([p['name'], p['quantity'], p['revenue']])

        t = Table(data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.purple),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(t)

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_product_report(product, calculation):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph(f"Fiche de Coût: {product['name']}", styles['Title']))

        data = [
            ["Élément", "Coût (DA)"],
            ["Ingrédients (Raw)", calculation['ingredient_cost_raw']],
            ["Ingrédients (+5% perte)", calculation['ingredient_cost_with_loss']],
            ["Main d'œuvre", calculation['labor_cost']],
            ["Packaging", calculation['packaging_cost']],
            ["COÛT VARIABLE TOTAL", calculation['total_variable_cost']]
        ]

        t = Table(data, colWidths=[200, 100])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold')
        ]))
        elements.append(t)

        doc.build(elements)
        buffer.seek(0)
        return buffer
